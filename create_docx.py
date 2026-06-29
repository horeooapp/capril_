import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX

doc = docx.Document()

# Styles
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

# Title
heading = doc.add_heading("RÉPONSES À L'AUDIT DU DÉVELOPPEMENT QAPRIL", 0)
heading.style.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)

doc.add_paragraph("Référence : QAPRIL/AUDIT/001/2026").bold = True
doc.add_paragraph("Destinataires : Frédéric Soumahoro · futurs Lead Developers").bold = True
doc.add_paragraph("Statut : Document de travail interne — Confidentiel").bold = True
doc.add_paragraph("Date : Juin 2026").bold = True

doc.add_heading("Partie 1 — Questionnaire au développeur", level=1)

doc.add_heading("1.1 Confidentialité et secrets", level=2)

q1 = doc.add_paragraph("Q1. L'outil a-t-il été utilisé en mode cloud ou en mode local (Ollama / modèle local) ?")
q1.style.font.bold = True

p1 = doc.add_paragraph("L'outil agentique (Google Antigravity) a été utilisé en ")
p1.add_run("mode cloud").bold = True
p1.add_run(". Le code source, les fichiers de configuration (y compris le fichier .env) et les requêtes ont été traités via les serveurs Cloud de Google DeepMind. L'agent tourne localement dans l'IDE mais le LLM qui l'anime est hébergé dans le cloud.")

q2 = doc.add_paragraph("Q2. Des clés ou secrets RÉELS de production ont-ils été utilisés pendant le développement ?")
q2.style.font.bold = True

p2 = doc.add_paragraph()
alert_run = p2.add_run("OUI. ")
alert_run.bold = True
alert_run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
p2.add_run("Après vérification de l'environnement local, le fichier .env contient des identifiants et clés d'API bien réels qui ont été exposés au modèle Cloud :")

doc.add_paragraph("Le mot de passe de la base de données PostgreSQL (QaprilSecure2026!).", style='List Bullet')
doc.add_paragraph("La clé de production pour le service d'e-mail Resend (re_YRdesQfm_...).", style='List Bullet')
doc.add_paragraph("Le mot de passe du serveur SMTP Hostinger (Anny_jay206).", style='List Bullet')
doc.add_paragraph("Une clé d'API Anthropic complète et fonctionnelle.", style='List Bullet')
doc.add_paragraph("Un token JWT et API complet pour le service Twenty.", style='List Bullet')

p2b = doc.add_paragraph("Toutes ces clés doivent être considérées comme compromises par précaution et faire l'objet d'une rotation immédiate.")

q3 = doc.add_paragraph("Q3. Les secrets sont-ils stockés en variables d'environnement, ou y a-t-il des secrets « en dur » dans le code ?")
q3.style.font.bold = True
doc.add_paragraph("Les secrets identifiés sont stockés de façon conventionnelle dans le fichier .env et appelés via process.env. L'audit n'a pas révélé de mots de passe de base de données ou de clés API critiques insérés \"en dur\" dans le code TypeScript/JavaScript. Cependant, certaines clés (comme CinetPay) semblent gérées de manière partielle. L'exposition principale vient du fichier .env lui-même qui a été analysé par l'agent.")

q4 = doc.add_paragraph("Q4. Le dépôt de code (GitHub) a-t-il toujours été privé ? Qui y a / y a eu accès ?")
q4.style.font.bold = True
doc.add_paragraph("Le dépôt utilisé est horeooapp/capril_. Bien qu'il s'agisse probablement d'un dépôt privé au niveau de l'organisation GitHub, l'historique Git devra être audité rigoureusement pour s'assurer que le fichier .env n'a jamais été commité par inadvertance. Si c'est le cas, un nettoyage profond (BFG Repo-Cleaner) ou la rotation des clés est impérative.")

doc.add_heading("1.2 Architecture et maîtrise", level=2)

q5 = doc.add_paragraph("Q5. Peux-tu expliquer l'architecture générale : modules, base de données, flux d'authentification, gestion des paiements ?")
q5.style.font.bold = True

p5_1 = doc.add_paragraph("", style='List Bullet')
p5_1.add_run("Frontend/Backend Web : ").bold = True
p5_1.add_run("L'application utilise Next.js (App Router) englobant à la fois le site portail et les API backend.")

p5_2 = doc.add_paragraph("", style='List Bullet')
p5_2.add_run("Mobile : ").bold = True
p5_2.add_run("Deux applications natives distinctes (Locataire et Propriétaire) développées en React Native (Expo).")

p5_3 = doc.add_paragraph("", style='List Bullet')
p5_3.add_run("Base de données : ").bold = True
p5_3.add_run("PostgreSQL hébergé, avec Prisma comme ORM pour garantir le typage fort et les migrations.")

p5_4 = doc.add_paragraph("", style='List Bullet')
p5_4.add_run("Authentification : ").bold = True
p5_4.add_run("NextAuth.js (Auth.js) pour la partie Web. Tokens JWT personnalisés pour la partie mobile.")

p5_5 = doc.add_paragraph("", style='List Bullet')
p5_5.add_run("Paiements : ").bold = True
p5_5.add_run("CinetPay intégré via des webhooks. Les transactions créent d'abord un PaymentIntent.")

q6 = doc.add_paragraph("Q6. Le code est-il documenté (commentaires, README, schéma de base de données à jour) ?")
q6.style.font.bold = True
doc.add_paragraph("Le schéma de la base de données (schema.prisma) est central, à jour et bien documenté. Le code contient des commentaires fonctionnels. En revanche, la documentation globale pour l'onboarding de l'équipe (ex: README complet) reste sommaire. Des rapports d'exécution (walkthroughs) locaux existent néanmoins dans les logs de l'agent.")

q7 = doc.add_paragraph("Q7. Le schéma de la base de données (schema.prisma) correspond-il exactement à ce qui tourne en production ?")
q7.style.font.bold = True
doc.add_paragraph("Oui. Le projet utilise l'approche \"Prisma First\". Toute modification de la base de données passe par une migration Prisma. L'ORM est la source unique de vérité, empêchant un décalage entre le code et la base.")

q8 = doc.add_paragraph("Q8. Quelles parties ont été écrites/modifiées à la main, et lesquelles ont été générées par l'agent sans relecture ligne à ligne ?")
q8.style.font.bold = True
doc.add_paragraph("Une grande proportion du code récent (notamment les API mobiles partagées, le flux complet de signature native OTP pour les baux, l'intégration CinetPay webhooks) a été générée itérativement par l'agent (Google Antigravity) sous supervision. Si la structure a été validée fonctionnellement, la relecture \"ligne à ligne\" par un développeur humain senior n'a pas nécessairement été systématique sur le code applicatif lourd.")

doc.add_heading("1.3 Tests et sécurité", level=2)

q9 = doc.add_paragraph("Q9. Existe-t-il des tests automatisés ? Que couvrent-ils réellement ?")
q9.style.font.bold = True
p9 = doc.add_paragraph()
r9 = p9.add_run("Insuffisant. ")
r9.font.color.rgb = RGBColor(0xE6, 0x7E, 0x22)
r9.bold = True
p9.add_run("Il existe des scripts de test unitaires isolés (ex: test-db.js), mais il manque une véritable suite de tests automatisés de bout en bout (Jest, Cypress). Le bon fonctionnement s'est surtout appuyé sur des tests manuels ou de l'auto-vérification.")

q10 = doc.add_paragraph("Q10. Comment sont gérées les règles métier critiques du CDC (confidentialité bailleur, mandats, SHA-256, webhooks HMAC) ?")
q10.style.font.bold = True
doc.add_paragraph("Ces règles sont gérées exclusivement au niveau du code métier (APIs Next.js). L'audit n'a pas révélé la présence de \"Triggers PostgreSQL\" natifs pour bloquer ces règles de sécurité directement dans le moteur de la base de données. Cela signifie que la sécurité dépend de l'absence de failles dans le code backend.")

q11 = doc.add_paragraph("Q11. Les webhooks de paiement vérifient-ils bien la signature HMAC-SHA256 sur le corps brut de la requête ?")
q11.style.font.bold = True
doc.add_paragraph("Dans le fichier route.ts de CinetPay, le webhook fonctionne sur un modèle Active Verification.\nAu lieu de vérifier un HMAC localement sur le corps brut, le système récupère l'identifiant de la transaction (cpm_trans_id) reçu via le webhook, puis effectue une requête HTTP (Call-back) à l'API de CinetPay (cinetpay.verifyPayment) pour obtenir le statut certifié de la transaction. Bien que cette méthode soit reconnue et sécurisée (elle empêche le spoofing car elle interroge la source), elle diffère de la vérification pure par signature HMAC hors-ligne. Un système d'idempotence a bien été repéré pour empêcher le double-crédit.")

doc.add_heading("Partie 2 — Bilan de la Checklist d'audit de code (Résumé)", level=1)

table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Niveau'
hdr_cells[1].text = 'Point de contrôle'
hdr_cells[2].text = 'Statut observé dans la base QAPRIL'

row_cells = table.add_row().cells
row_cells[0].text = '[CRITIQUE]'
row_cells[1].text = 'Aucun secret en dur dans le code'
row_cells[2].text = 'VALIDE (Utilisation de process.env)'

row_cells = table.add_row().cells
row_cells[0].text = '[CRITIQUE]'
row_cells[1].text = 'Tous les secrets en variables d\'environnement, jamais committés'
row_cells[2].text = 'À VÉRIFIER (Historique Git à vérifier. Le .env a été exposé au Cloud LLM)'

row_cells = table.add_row().cells
row_cells[0].text = '[CRITIQUE]'
row_cells[1].text = 'Rotation de TOUS les secrets effectuée si exposition possible'
row_cells[2].text = 'ÉCHEC (Les clés réelles sont toujours actives et dans le .env local. À faire en priorité absolue).'

row_cells = table.add_row().cells
row_cells[0].text = '[CRITIQUE]'
row_cells[1].text = 'Idempotence des paiements'
row_cells[2].text = 'VALIDE (Le webhook vérifie si le statut est déjà SUCCESS avant de processer).'

row_cells = table.add_row().cells
row_cells[0].text = '[IMPORTANT]'
row_cells[1].text = 'Triggers PostgreSQL des règles critiques présents'
row_cells[2].text = 'ÉCHEC (Pas de triggers SQL détectés. Logique gérée uniquement en backend).'

row_cells = table.add_row().cells
row_cells[0].text = '[IMPORTANT]'
row_cells[1].text = 'Tests automatisés présents sur les flux critiques'
row_cells[2].text = 'ÉCHEC (Aucune suite Jest/CI n\'est implémentée formellement pour l\'applicatif).'

doc.save('QAPRIL_Reponses_Audit_Final.docx')
